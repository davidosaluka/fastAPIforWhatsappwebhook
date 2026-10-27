from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# Organized Data categorized logically with Group headers followed immediately by their child variants
# Truncated prices from Meta catalog exports have also been completed.
data = [
    # --- RICE & MEALS ---
    ["Rice Meals", "Jollof", "mg_jollof", "Eligible", "-", "In stock", "NGN1,300.00"],
    ["Rice Meals", "Fried Rice", "mg_fried_rice", "Eligible", "-", "In stock", "NGN1,300.00"],
    ["Rice Meals", "Fried and Jollof", "mg_fried_jollof", "Eligible", "-", "In stock", "NGN1,300.00"],
    ["Rice Meals", "White rice", "mg_white_rice", "Eligible", "-", "In stock", "NGN1,300.00"],
    ["Rice Meals", "White Rice and beans", "mg_white_rice_beans", "Eligible", "-", "In stock", "NGN1,300.00"],

    # --- SWALLOWS & SOUPS ---
    ["Swallows & Soups", "Semo", "mg_semo", "Eligible", "-", "In stock", "NGN700.00"],
    ["Swallows & Soups", "Fufu", "mg_fufu", "Eligible", "-", "In stock", "NGN700.00"],
    ["Swallows & Soups", "Vegetable (Efo Riro)", "mg_veggies_efo", "Eligible", "-", "In stock", "NGN500.00"],
    ["Swallows & Soups", "Egusi", "mg_egusi", "Eligible", "-", "In stock", "NGN200.00"],
    ["Swallows & Soups", "Ogbono", "mg_ogbono", "Eligible", "-", "In stock", "NGN200.00"],

    # --- PROTEIN: CHICKEN (GROUP + 3 VARIANTS) ---
    ["Chicken", "Chicken (Group)", "4cqemme5ey", "Eligible", "3", "In stock (3 variants)", "NGN3,000.00 - NGN4,000.00"],
    ["Chicken", "  ↳ Small Chicken", "mg_chicken_small", "Eligible", "-", "In stock", "NGN3,000.00"],
    ["Chicken", "  ↳ Medium Chicken", "mg_chicken_medium", "Eligible", "-", "In stock", "NGN3,500.00"],
    ["Chicken", "  ↳ Big Chicken", "mg_chicken_big", "Eligible", "-", "In stock", "NGN4,000.00"],

    # --- PROTEIN: SMALL BEEF (GROUP + 6 VARIANTS) ---
    ["Small Beef", "Small Beef (Group)", "-", "Eligible", "6", "In stock (6 variants)", "NGN200.00 - NGN1,200.00"],
    ["Small Beef", "  ↳ Small Beef One Piece", "mg_beef_one", "Eligible", "-", "In stock", "NGN200.00"],
    ["Small Beef", "  ↳ Small Beef two Pieces", "mg_small_beef_two", "Eligible", "-", "In stock", "NGN400.00"],
    ["Small Beef", "  ↳ Small Beef three Pieces", "mg_small_beef_three", "Eligible", "-", "In stock", "NGN600.00"],
    ["Small Beef", "  ↳ Small Beef four Pieces", "mg_small_beef_four", "Eligible", "-", "In stock", "NGN800.00"],
    ["Small Beef", "  ↳ Small Beef five Pieces", "mg_small_beef_five", "Eligible", "-", "In stock", "NGN1,000.00"],
    ["Small Beef", "  ↳ Small Beef six Pieces", "mg_small_beef_six", "Eligible", "-", "In stock", "NGN1,200.00"],

    # --- PROTEIN: BIG BEEF (GROUP + 6 VARIANTS) ---
    ["Big Beef", "Big Beef (Group)", "me40n6zunu", "Eligible", "6", "In stock (6 variants)", "NGN500.00 - NGN3,000.00"],
    ["Big Beef", "  ↳ Big beef One Piece", "mg_big_beef_one", "Eligible", "-", "In stock", "NGN500.00"],
    ["Big Beef", "  ↳ Big beef two Pieces", "mg_big_beef_two", "Eligible", "-", "In stock", "NGN1,000.00"],
    ["Big Beef", "  ↳ Big beef three Pieces", "mg_big_beef_three", "Eligible", "-", "In stock", "NGN1,500.00"],
    ["Big Beef", "  ↳ Big beef four Pieces", "mg_big_beef_four", "Eligible", "-", "In stock", "NGN2,000.00"],
    ["Big Beef", "  ↳ Big beef five Pieces", "mg_big_beef_five", "Eligible", "-", "In stock", "NGN2,500.00"],
    ["Big Beef", "  ↳ Big beef six Pieces", "mg_big_beef_six", "Eligible", "-", "In stock", "NGN3,000.00"],

    # --- PROTEIN: GOAT MEAT (GROUP + 2 VARIANTS) ---
    ["Goat Meat", "Goat Meat & Pepper Soup (Group)", "0f6b5aesff", "Eligible", "2", "In stock (2 variants)", "NGN500.00 - NGN2,000.00"],
    ["Goat Meat", "  ↳ Goat meat", "mg_goat_meat", "Eligible", "-", "In stock", "NGN500.00"],
    ["Goat Meat", "  ↳ Goat meat pepper soup", "mg_goat_meat_soup", "Eligible", "-", "In stock", "NGN2,000.00"],

    # --- PROTEIN: EGGS (GROUP + 2 VARIANTS) ---
    ["Eggs", "Eggs (Group)", "pfkkerqi57", "Eligible", "2", "In stock (2 variants)", "NGN300.00 - NGN1,000.00"],
    ["Eggs", "  ↳ Egg", "mg_egg", "Eligible", "-", "In stock", "NGN300.00"],
    ["Eggs", "  ↳ Fried egg", "mg_fried_egg", "Eligible", "-", "In stock", "NGN1,000.00"],

    # --- SIDES & EXTRAS (PLANTAIN, MOI MOI + SALAD) ---
    ["Sides & Extras", "Plantain (Group)", "8h5fh5bhx0", "Eligible", "2", "In stock (2 variants)", "NGN500.00 - NGN1,000.00"],
    ["Sides & Extras", "  ↳ Plantain (Single serving)", "mg_plantain", "Eligible", "-", "In stock", "NGN500.00"],
    ["Sides & Extras", "  ↳ Plantain two serving", "mg_plantain_two", "Eligible", "-", "In stock", "NGN1,000.00"],
    ["Sides & Extras", "Moi Moi (Group)", "tmvcd4wq76", "Eligible", "2", "In stock (2 variants)", "NGN1,000.00 - NGN2,000.00"],
    ["Sides & Extras", "  ↳ Moi Moi", "mg_moi_moi", "Eligible", "-", "In sto`ck", "NGN1,000.00"],
    ["Sides & Extras", "  ↳ Moi Moi two serving", "mg_moi_moi_two", "Eligible", "-", "In stock", "NGN2,000.00"],
    ["Sides & Extras", "Salad", "mg_salad", "Eligible", "-", "In stock", "NGN500.00"],

    # --- DRINKS ---
    ["Drinks", "Water", "mg_mr_v", "Eligible", "-", "In stock", "NGN300.00"],
    ["Drinks", "Coke", "mg_coke", "Eligible", "-", "In stock", "NGN500.00"],
    ["Drinks", "Maltina", "mg_maltina", "Eligible", "-", "In stock", "NGN700.00"],
]

columns = ["Category", "Name", "Content ID / Group ID", "Status", "Variants", "Availability", "Price"]
file_name = "food_products.xlsx"

# 1. Initialize Workbook & Sheet
wb = Workbook()
ws = wb.active
ws.title = "Products"

# Ensure grid lines are visible
ws.views.sheetView[0].showGridLines = True

# 2. Append Headers and Data
ws.append(columns)
for row in data:
    ws.append(row)

# 3. Create & Apply Styled Table
tab = Table(displayName="FoodProductsTable", ref=f"A1:G{len(data)+1}")
style = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)
tab.tableStyleInfo = style
ws.add_table(tab)

# 4. Styling enhancements:
# Highlight group headers with bold text and subtle highlight fill
group_font = Font(name="Calibri", size=11, bold=True)
group_fill = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid")
normal_font = Font(name="Calibri", size=11)
variant_font = Font(name="Calibri", size=11, color="333333")

for row_idx in range(2, len(data) + 2):
    row_data = data[row_idx - 2]
    is_group = "(Group)" in row_data[1]
    is_child = row_data[1].startswith("  ↳")
    
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if is_group:
            cell.font = group_font
            cell.fill = group_fill
        elif is_child:
            cell.font = variant_font
        else:
            cell.font = normal_font

# 5. Auto-fit Column Widths for readability
for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

# 6. Save the Workbook
try:
    wb.save(file_name)
    print(f"Successfully generated organized {file_name}")
except PermissionError:
    alt_name = "food_products_organized.xlsx"
    wb.save(alt_name)
    print(f"Note: '{file_name}' is currently open in Excel. Saved to '{alt_name}' instead.")
